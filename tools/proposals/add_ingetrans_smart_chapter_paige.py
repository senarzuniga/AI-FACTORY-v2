"""Build the Paige proposal revision that adds the INGETRANS Smart chapter after Block A.

Reads the re2 proposal, inserts a new Heading 1 chapter (Heading 2 subsections, List Bullet
paragraphs and tables cloned from the document's own Block A table) immediately before the
Block B heading, adds one executive-summary bullet and one price-overview note, and writes a
new file. The original re2 file is never modified.
"""

from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

BASE = Path(r"C:\Users\isena\Documents\INGECART\COMMERCIAL\PROYECTOS\Sterner Global\Paige")
SOURCE = BASE / "NewCorr_Plant_CORRUGATOR_AREA_INGECART_PROPOSAL_2026-08-26 re2.docx"
TARGET = BASE / "NewCorr_Plant_CORRUGATOR_AREA_INGECART_PROPOSAL_2026-09-14 re3 INGETRANS SMART.docx"

BLOCK_B_HEADING_PREFIX = "6  BLOCK B"
TEMPLATE_TABLE_INDEX = 5  # "Parameter | Proposal basis" table of Block A


def find_paragraph(doc, prefix: str, style: str = "Heading 1") -> Paragraph:
    for p in doc.paragraphs:
        if p.style.name == style and p.text.strip().startswith(prefix):
            return p
    raise SystemExit(f"Anchor paragraph not found: {prefix} [{style}]")


def clear_paragraph(p: Paragraph) -> None:
    # Remove every child except paragraph properties so hyperlinks/field runs in TOC clones vanish too.
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


class Inserter:
    """Inserts block-level elements before an anchor paragraph, preserving document order."""

    def __init__(self, doc, anchor: Paragraph, template_table: Table):
        self.doc = doc
        self.anchor = anchor._p
        self.template_table = template_table

    def paragraph(self, text: str = "", style: str = "Normal", bold: bool = False) -> Paragraph:
        p = self.doc.add_paragraph(style=style)
        if text:
            p.add_run(text).bold = bold
        self.anchor.addprevious(p._p)
        return p

    def heading(self, text: str, level: int) -> Paragraph:
        return self.paragraph(text, style=f"Heading {level}")

    def bullets(self, items) -> None:
        for item in items:
            self.paragraph(item, style="List Bullet")

    def table(self, rows, widths=None) -> Table:
        cols = len(rows[0])
        total_dxa = 9354  # usable width of the source document tables
        if widths is None:
            widths = [1.0] * cols
        scale = total_dxa / sum(widths)
        col_dxa = [int(w * scale) for w in widths]

        tbl = copy.deepcopy(self.template_table._tbl)
        for tr in tbl.findall(qn("w:tr")):
            tbl.remove(tr)
        grid = tbl.find(qn("w:tblGrid"))
        grid_template = grid.find(qn("w:gridCol"))
        for gc in grid.findall(qn("w:gridCol")):
            grid.remove(gc)
        for width in col_dxa:
            gc = copy.deepcopy(grid_template)
            gc.set(qn("w:w"), str(width))
            grid.append(gc)

        template_rows = self.template_table.rows
        for r_idx in range(len(rows)):
            src = template_rows[0] if r_idx == 0 else template_rows[1]
            tr = copy.deepcopy(src._tr)
            tcs = tr.findall(qn("w:tc"))
            while len(tcs) > cols:
                tr.remove(tcs.pop())
            while len(tcs) < cols:
                tr.append(copy.deepcopy(tcs[-1]))
                tcs = tr.findall(qn("w:tc"))
            for tc, width in zip(tcs, col_dxa):
                tcw = tc.find(qn("w:tcPr")).find(qn("w:tcW"))
                if tcw is not None:
                    tcw.set(qn("w:w"), str(width))
                    tcw.set(qn("w:type"), "dxa")
            tbl.append(tr)

        self.anchor.addprevious(tbl)
        table = Table(tbl, self.anchor.getparent())
        for r_idx, values in enumerate(rows):
            for c_idx, value in enumerate(values):
                cell = table.cell(r_idx, c_idx)
                for extra in cell.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
                para = cell.paragraphs[0]
                for r in list(para.runs):
                    r._r.getparent().remove(r._r)
                run = para.add_run(str(value))
                run.bold = r_idx == 0
                run.font.size = Pt(9)
                if r_idx == 0:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        self.paragraph("")
        return table


def build_chapter(ins: Inserter) -> None:
    ins.heading("5A  BLOCK A-SMART - INGETRANS SMART RAIL-LESS OPTION", 1)
    ins.paragraph("")
    ins.paragraph("INGETRANS Smart - rail-less and battery-powered variant of Block A", bold=True)
    ins.paragraph(
        "INGETRANS Smart delivers the same automatic reel delivery and return function described in Block A, "
        "but replaces the two elements of INGETRANS that generate most of the civil work and fixed infrastructure: "
        "the embedded steel running rails of the transfer carriage and the conductor line (catenary, Vahle type) "
        "that powers the carriage while it moves. In INGETRANS Smart the carriage runs directly on the plant floor "
        "on modular drive and steering wheels (JNOV TECH MDS/DDW technology), is guided by a floor DataMatrix line "
        "read by optical sensors, and is powered by on-board lithium-ion batteries recharged at a parking position "
        "between missions."
    )
    ins.paragraph(
        "The reel pick-up and release mechanism, the exchange area with its ramps, the ten roll-stand tracks, the "
        "control philosophy, the MES/ERP interface and the safety architecture of Block A are retained. Block A-Smart "
        "is offered as an alternative to Block A; the customer selects one of the two configurations."
    )

    ins.heading("5A.1 What is removed and what replaces it", 2)
    ins.table([
        ["Block A element", "Function", "INGETRANS Smart"],
        ["Embedded steel running rails along the more than 60 m travel corridor", "Guide and support the transfer carriage", "Removed. Polyurethane drive wheels run on the existing concrete floor; guidance by DataMatrix floor line (+/-0.1 mm) read by optical sensors."],
        ["Conductor line (Vahle-type catenary), supports and current collector", "Power the carriage in motion", "Removed. On-board lithium-ion batteries (one per drive wheel), control cabinet with integrated chargers, parking/charging position with 230 V IEC socket."],
        ["Rail civil works: trenching, embedded plates, levelling, anchoring, floor repair", "Rail foundation", "Removed. Replaced by a floor flatness survey (DIN 18202 table 3 line 3) and local floor repair only where required."],
        ["Mechanical end stops and rail-mounted limit switches", "Travel limits", "Replaced by DataMatrix position codes, software zones and safety scanners."],
        ["Single traction motor and gearbox running on rails", "Carriage traction", "Replaced by 2 x DDW-4TP drive and steering wheels (four brushless motors, differential steering, omnidirectional motion)."],
        ["Reel pick-up/lifting mechanism, exchange area, ramps, ten roll-stand tracks", "Reel handling and delivery", "Retained without functional change."],
    ], widths=[2.3, 1.7, 4.0])

    ins.heading("5A.2 Functional scope", 2)
    ins.bullets([
        "Automatic pickup of validated reels from the exchange position, delivery to the selected roll-stand track and return of unused or partially consumed reels, identical to Block A.",
        "Rail-less travel along the corrugator corridor on modular drive and steering wheels; the route is defined by a DataMatrix floor line that can be re-laid without civil works if the layout changes.",
        "Three operating levels: Smart Manual (radio remote control for commissioning, maintenance and degraded mode), Smart Guided (automatic line following with position codes at each roll stand) and Smart Orchestrated (INGECART master PLC commands speed, direction and missions over EtherNet/IP and supervises state, battery level and faults).",
        "Battery management integrated in the INGECART sequence: opportunistic charging between missions at the parking position; missions are not released below the configured battery threshold and the condition is reported.",
        "Central HMI, sequence control, alarms, diagnostics, movement history and MES/ERP interface as in Block A.",
        "Safety interlocks between carriage, fixed tracks, charging position and protected operating zones.",
        "Factory testing of the carriage with drive wheels and control cabinet before shipment.",
    ])

    ins.heading("5A.3 Technical basis", 2)
    ins.table([
        ["Parameter", "Block A (rail-guided)", "Block A-Smart (rail-less) - proposal basis"],
        ["System type", "Surface rail-guided transfer carriage", "Rail-less carriage on 2 x JNOV DDW-4TP drive and steering wheels plus curved auxiliary castors; DataMatrix line guidance."],
        ["Transfer speed", "Reference up to 100-80 m/min", "Reference up to 2 km/h (33 m/min) contractual for DDW-4TP; range value up to 2.5 km/h. Reduced automatically in safe-speed zones near people and stations."],
        ["Positioning accuracy", "Encoder / rail end stops", "+/-0.1 mm on DataMatrix line; +/-1 mm on painted line."],
        ["Acceleration / deceleration", "Approx. 1.5 s", "Configurable in the drive cabinet; reference 1.5-4.5 s ramps."],
        ["Pickup / drop-off", "Approx. 6 s per interface", "Approx. 6 s per interface (mechanism retained)."],
        ["Reel envelope", "Up to 1,500 mm diameter, up to 2,800 mm width, below 3,500 kg", "Same envelope; carriage mass distribution engineered so that the drive wheels carry more than 50% of the total mass and castors stay below 70% of their rating."],
        ["Drive wheel capacity", "-", "4,000 kg carrying and 4,000 kg towed per DDW; 1.6 kN nominal / 4.5 kN maximum traction; polyurethane tyre diameter 250 mm; 60 kg per wheel; IP54."],
        ["Power supply", "Conductor line (Vahle type)", "Lithium-ion batteries, one per drive wheel, 600 charge cycles; charging at parking position via 230 V IEC C14 socket."],
        ["Floor requirements", "Rail foundation", "Concrete or asphalt floor; flatness 4 mm/1 m and 10 mm/4 m (DIN 18202); slope up to 3%; steps up to 12 mm; ground contact pressure above 10 N/mm2."],
        ["Controls", "PLC/HMI, PROFINET", "INGECART PLC/HMI as master; EtherNet/IP link to the drive cabinet (PROFINET on request); 2.4 GHz DSSS radio remote for manual mode."],
        ["Safety", "Safety PLC, scanners, interlocks, E-stops", "Drive system SS1 stop category 1 (PLd), STO on all drives, power-off safety brakes, E-stop PLd/Cat 3 on the remote, stop on radio loss; INGECART adds area scanners with warning and protective fields, cabinet E-stop, safety auxiliary inputs and zone interlocks."],
        ["Environment", "-", "0 to +40 C, relative humidity up to 85%, indoor."],
        ["Drive wheel service life", "-", "2,500 km or 10,000 operating hours nominal; tyre rims replaceable."],
    ], widths=[1.8, 2.4, 4.4])

    ins.heading("5A.4 Capacity and speed assessment", 2)
    ins.paragraph(
        "INGETRANS Smart travels slower than the rail-guided carriage. The table compares one-way travel times "
        "for the Paige corridor using trapezoidal speed profiles and the same 6 s pick-up and 6 s drop-off "
        "interfaces. The reference demand of the corrugator model used in Block A is 8.8 reels/hour."
    )
    ins.table([
        ["Distance exchange area to roll stand", "Block A one-way (80 m/min)", "Block A-Smart one-way (2 km/h)", "Block A-Smart one-way (2.5 km/h)"],
        ["12 m (nearest roll stand)", "approx. 14 s", "approx. 26 s", "approx. 22 s"],
        ["30 m (mid corridor)", "approx. 27 s", "approx. 59 s", "approx. 48 s"],
        ["60 m (far end of corridor)", "approx. 50 s", "approx. 113 s", "approx. 91 s"],
        ["Complete mission, 30 m round trip incl. pick-up and drop-off", "approx. 66 s (approx. 54 missions/h)", "approx. 129 s (approx. 28 missions/h)", "approx. 108 s (approx. 33 missions/h)"],
    ], widths=[2.6, 2.0, 2.2, 2.2])
    ins.bullets([
        "At 8.8 reels/hour, INGETRANS Smart keeps carriage utilisation below 35% even with 30 m average missions; dual-reel transport and simultaneous return of partial reels further reduce the number of missions.",
        "For the far roll stands (about 60 m) the one-way time approaches two minutes at 2 km/h; the INGECART sequence compensates with anticipated reel staging and dual-reel trips. If the final production mix requires more than about 20 reels/hour with long average travel, the rail-guided Block A or a second Smart carriage is recommended.",
        "The final capacity statement will be confirmed with the INGECART digital-twin model on the approved Paige layout during engineering, using the same methodology as the Block A benchmark.",
    ])

    ins.heading("5A.5 Civil works, installation and maintenance impact", 2)
    ins.table([
        ["Aspect", "Block A (rail-guided)", "Block A-Smart (rail-less)"],
        ["Civil works", "Trenches, embedded plates and levelling for more than 60 m of steel running rails; supports for the conductor line", "No rail trenches or embedded plates; floor flatness survey and local repair only. Floor line applied on the existing surface."],
        ["Fixed infrastructure supplied", "Rails, conductor line, current collector, end stops", "DataMatrix floor line, parking/charging position, safety scanners on the carriage"],
        ["Site services", "Rail installation block (14 days, 1 technician) plus mechanical assembly", "Rail installation block removed; mechanical assembly limited to carriage, exchange area and tracks. Final day count confirmed after layout approval."],
        ["Production disruption during installation", "Rail works inside the corrugator corridor", "Minimal: floor line marking and carriage positioning"],
        ["Future layout changes", "New rail sections and civil works", "Re-lay the floor line and update position codes in software"],
        ["Preventive maintenance", "Rail wear and alignment, conductor line and collector brushes, traction gearbox", "Tyre wear inspection and rim replacement, battery health and charge cycles, optical sensor cleaning, floor line condition"],
        ["Spare parts", "Collector brushes, rail fixings, traction motor", "Drive wheel exchangeable as a unit through a single M23 hybrid connector; batteries; optical sensor"],
        ["Single points of failure", "Conductor line and collector", "None on power supply; drive wheels interchangeable; manual remote mode available"],
        ["Housekeeping requirements", "Rail groove cleaning", "Floor line kept free of paper dust and forklift damage; protected tape option"],
    ], widths=[1.8, 3.3, 3.5])

    ins.heading("5A.6 Safety basis for the rail-less carriage", 2)
    ins.bullets([
        "Drive system safety functions certified by the wheel manufacturer: Safety Stop 1 (stop category 1 according to IEC 60204-1) at PLd, STO on all drives, power-off safety brakes acting also as parking brakes, emergency stop PLd/Cat 3 on the radio remote and automatic SS1 on loss of radio signal.",
        "INGECART completes the machine: laser area scanners with warning and protective fields in both travel directions, emergency stop on the carriage cabinet, dead-man handle for manual mode, up to two additional safety sensor inputs, safe-speed zones at roll stands and exchange area, and interlocks with tracks and charging position.",
        "The drive system is delivered as partly completed machinery (Machinery Directive 2006/42/EC, EMC Directive 2014/30/EU). The risk assessment and CE marking of the complete INGETRANS Smart carriage are part of the INGECART scope; ISO 3691-4 requirements for driverless industrial trucks are applied when the carriage operates without an operator in the zone.",
        "Radio environment: one IEEE 802.15.4 channel (2.4 GHz, channels 11-25) with noise below -80 dBm must be available; a Wi-Fi coexistence check is included in engineering.",
    ])

    ins.heading("5A.7 Operational value versus Block A", 2)
    ins.bullets([
        "Same automatic reel logistics, same MES/ERP integration and same forklift-free corrugator area as Block A.",
        "Elimination of more than 60 m of embedded rails and of the conductor line: less civil work, shorter and less disruptive installation, no rail or catenary maintenance.",
        "Layout flexibility: additional roll stands, a relocated exchange area or a second corrugator can be served by re-laying the floor line instead of extending rails.",
        "No single point of failure in the power supply; interchangeable drive wheels and manual remote mode provide a controlled degraded mode.",
        "Battery energy use is measurable per mission and reported to the plant systems, supporting the same request-to-delivery and starvation KPIs used in Block A.",
    ])

    ins.heading("5A.8 Site prerequisites for Block A-Smart", 2)
    ins.bullets([
        "Concrete or asphalt floor in the corridor meeting DIN 18202 table 3 line 3 flatness (4 mm/1 m, 10 mm/4 m), slope up to 3%, steps up to 12 mm and ground contact pressure above 10 N/mm2. Floor survey performed by INGECART during engineering; repairs are customer scope.",
        "Free corridor for the DataMatrix floor line with a minimum curve radius of 0.5 m and protection against forklift traffic.",
        "Parking/charging position with 230 V single-phase supply and access for battery exchange.",
        "One free 2.4 GHz IEEE 802.15.4 channel and line of sight between remote and carriage in manual mode.",
        "Ambient conditions 0 to +40 C and relative humidity up to 85%.",
    ])

    ins.heading("5A.9 Engineering qualification and price basis", 2)
    ins.bullets([
        "Speed, positioning, capacity and safety values of the drive system are manufacturer data for the DDW-4TP wheel and will be confirmed in the Functional Design Specification and the FAT/SAT acceptance matrix.",
        "Battery autonomy per charge will be stated in the FDS from the manufacturer's load-dependent autonomy curve and validated during SAT with the Paige reel matrix and mission profile.",
        "The Smart carriage frame, mass distribution and wheel loading are engineered by INGECART; the reel envelope of Block A remains the governing design basis.",
        "Block A-Smart supply is offered at the same EXW price as Block A. The rail installation service block (14 days, 1 technician) is not required for Block A-Smart and is removed from the site-services table when this option is selected; the mechanical assembly block is re-quoted after layout approval. Drive-system lead time (approximately 24 weeks from order) is included in the project schedule.",
        "Block A and Block A-Smart are mutually exclusive alternatives within this proposal.",
    ])
    ins.paragraph("")


def insert_after(doc, anchor_prefix: str, text: str, style: str | None = None, new_style: str | None = None) -> bool:
    for p in doc.paragraphs:
        if p.text.startswith(anchor_prefix) and (style is None or p.style.name == style):
            new_p = copy.deepcopy(p._p)
            p._p.addnext(new_p)
            np_ = Paragraph(new_p, p._parent)
            clear_paragraph(np_)
            np_.add_run(text)
            if new_style:
                np_.style = doc.styles[new_style]
            return True
    return False


def main() -> int:
    if TARGET.exists():
        print(f"Target already exists, not overwriting: {TARGET}")
        return 1
    doc = Document(str(SOURCE))
    anchor = find_paragraph(doc, BLOCK_B_HEADING_PREFIX)
    build_chapter(Inserter(doc, anchor, doc.tables[TEMPLATE_TABLE_INDEX]))

    insert_after(
        doc,
        "5.5 Internal scenario benchmark",
        "5A  BLOCK A-SMART - INGETRANS SMART RAIL-LESS OPTION\t1",
        style="toc 2",
        new_style="toc 1",
    )
    insert_after(
        doc,
        "Automatic delivery and return of paper reels",
        "Optional Block A-Smart: the same INGETRANS reel logistics on a rail-less, battery-powered carriage "
        "(JNOV drive wheels, DataMatrix guidance), removing more than 60 m of embedded rails and the conductor "
        "line together with their civil works.",
    )
    insert_after(
        doc,
        "The RFID validation station price is included in Block C",
        "Block A-Smart (INGETRANS Smart, rail-less option described in chapter 5A) is offered at the same EXW "
        "supply price as Block A and replaces it when selected. When Block A-Smart is chosen, the rail installation "
        "service block is not applicable and the mechanical assembly block is re-quoted after layout approval.",
    )
    doc.save(str(TARGET))
    print(f"Saved: {TARGET}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
